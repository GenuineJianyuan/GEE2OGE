"""读取表结构 docx 文档"""
import docx

doc = docx.Document(r'd:\docs\交投\文档\全球院工作\geeToOGE\__pycache__\geeToOgeSke\表结构整理.docx')

print("=" * 80)
print("段落内容：")
print("=" * 80)
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")

print("\n" + "=" * 80)
print("表格内容：")
print("=" * 80)
for ti, t in enumerate(doc.tables):
    print(f"\n--- 表格 {ti} ---")
    for ri, row in enumerate(t.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"  行{ri}: {' | '.join(cells)}")
